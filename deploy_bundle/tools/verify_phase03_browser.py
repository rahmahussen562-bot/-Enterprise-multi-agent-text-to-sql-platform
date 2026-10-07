"""Run actual frontend/API/browser workflows with ephemeral private test identities."""
import asyncio
import json
import os
from pathlib import Path
import secrets
import subprocess
import sys
import time
root=Path(__file__).resolve().parents[1]
frontend=Path('D:/BIRD-Interact/frontend')
runtime=root/'.runtime'/'phase3-preview';runtime.mkdir(parents=True,exist_ok=True)
for name,folder in {'TEMP':'tmp','TMP':'tmp','TMPDIR':'tmp','npm_config_cache':'npm-cache','PLAYWRIGHT_BROWSERS_PATH':'browsers'}.items():
    path=frontend/'.runtime'/folder;path.mkdir(parents=True,exist_ok=True);os.environ[name]=str(path)
os.environ['PYTHONDONTWRITEBYTECODE']='1';os.environ['NODE_OPTIONS']='--max-old-space-size=768 --max-semi-space-size=2'
os.chdir(root);sys.path.insert(0,str(root))
stop=runtime/'stop'
if '--serve' in sys.argv:
    import uvicorn
    server=uvicorn.Server(uvicorn.Config('api.main:app',host='127.0.0.1',port=8000,ws_max_size=16384))
    async def serve():
        async def monitor():
            while not stop.exists():await asyncio.sleep(.1)
            server.should_exit=True
        task=asyncio.create_task(monitor())
        try:await server.serve()
        finally:task.cancel();await asyncio.gather(task,return_exceptions=True)
    asyncio.run(serve());raise SystemExit(0)
from core.auth import password_hash
import httpx
password=secrets.token_urlsafe(24)
accounts=[{'user_id':identifier,'role':role,'password_hash':password_hash(password)} for identifier,role in [('phase3_sales','sales_analyst'),('phase3_inventory','inventory_lead')]]
accounts_file=runtime/'accounts.json';accounts_file.write_text(json.dumps(accounts),encoding='utf-8');stop.unlink(missing_ok=True)
environment=dict(os.environ,SENTINEL_AUTH_FILE=str(accounts_file),SENTINEL_SESSION_KEY=secrets.token_urlsafe(48),SENTINEL_API_STATE=str(runtime/'jobs.sqlite'),
    SENTINEL_DEMO_MODE='1',SQLITE_PATH=str(root/'data'/'chinook.db'),MSSQL_TIMEOUT='1',OPENAI_API_KEY='',GEMINI_API_KEY='',ANTHROPIC_API_KEY='',
    SENTINEL_ALLOWED_ORIGINS='http://localhost:5173,http://127.0.0.1:5173',SENTINEL_E2E_PASSWORD=password)
node=frontend/'.runtime'/'node'/'node.exe'
flags=subprocess.CREATE_NO_WINDOW if os.name=='nt' else 0
with (runtime/'api.log').open('w',encoding='utf-8') as backend_log,(runtime/'vite.log').open('w',encoding='utf-8') as frontend_log:
    backend=subprocess.Popen([sys.executable,'-B',__file__,'--serve'],cwd=root,env=environment,stdout=backend_log,stderr=backend_log,creationflags=flags)
    vite=subprocess.Popen([str(node),str(frontend/'node_modules'/'vite'/'bin'/'vite.js'),'--host','127.0.0.1'],cwd=frontend,env=environment,stdout=frontend_log,stderr=frontend_log,creationflags=flags)
    try:
        with httpx.Client(timeout=2,trust_env=False) as client:
            for url in ('http://127.0.0.1:8000/api/v1/health','http://127.0.0.1:5173'):
                deadline=time.monotonic()+60
                while True:
                    try:client.get(url).raise_for_status();break
                    except httpx.HTTPError:
                        if time.monotonic()>deadline:raise RuntimeError('Preview startup timed out; inspect D-drive preview logs.')
                        time.sleep(.25)
        subprocess.run([str(node),'tools/browser-check.mjs'],cwd=frontend,env=environment,check=True)
        import openpyxl
        workbook=openpyxl.load_workbook(frontend/'.runtime'/'export-check.xlsx',data_only=False)
        sheet=workbook.active
        assert sheet.max_row==6 and sheet.max_column==5 and not any(cell.data_type=='f' for row in sheet for cell in row)
        print('Actual Excel export reopened: 5 rows, 5 columns, no formulas.')
        workbook.close()
    finally:
        stop.write_text('stop',encoding='ascii')
        backend.wait(timeout=25)
        vite.terminate();vite.wait(timeout=10)
